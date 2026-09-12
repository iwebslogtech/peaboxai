from __future__ import annotations
import base64
import io
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".py", ".ps1", ".js", ".ts", ".yaml", ".yml", ".log", ".ini", ".cfg"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
MAX_EXTRACTED_CHARS = 60000


def extension(name):
    return Path(name).suffix.lower()


def extract_uploaded_file(uploaded):
    suffix = extension(uploaded.name)
    raw = uploaded.getvalue()
    if suffix in TEXT_EXTENSIONS:
        return raw.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(raw))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(io.BytesIO(raw))
        output = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                output.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(output)
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        output = []
        for ws in wb.worksheets:
            output.append(f"### Worksheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                output.append("\t".join("" if value is None else str(value) for value in row))
        wb.close()
        return "\n".join(output)
    if suffix == ".pptx":
        prs = Presentation(io.BytesIO(raw))
        output = []
        for index, slide in enumerate(prs.slides, 1):
            output.append(f"### Slide {index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    output.append(shape.text)
        return "\n".join(output)
    raise ValueError(f"Unsupported attachment type: {suffix or 'unknown'}")


def build_user_content(prompt, uploads, supports_vision):
    text_sections = []
    images = []
    used = 0
    for uploaded in uploads or []:
        suffix = extension(uploaded.name)
        if suffix in IMAGE_EXTENSIONS:
            if not supports_vision:
                raise ValueError(f"The selected model is not detected as vision-capable: {uploaded.name}")
            media = "image/jpeg" if suffix in {".jpg", ".jpeg"} else f"image/{suffix[1:]}"
            images.append({"name": uploaded.name, "media_type": media, "base64": base64.b64encode(uploaded.getvalue()).decode("ascii")})
            continue
        extracted = extract_uploaded_file(uploaded)
        available = MAX_EXTRACTED_CHARS - used
        if available <= 0:
            text_sections.append("[Further attachment text omitted because the extraction limit was reached.]")
            break
        clipped = extracted[:available]
        used += len(clipped)
        text_sections.append(f"\n\n===== ATTACHMENT: {uploaded.name} =====\n{clipped}")
    combined = prompt
    if text_sections:
        combined += "\n\nUse the following extracted attachment content as source material:" + "".join(text_sections)
    return combined, images
